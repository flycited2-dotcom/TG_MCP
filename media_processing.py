"""Bounded local media/document readers. No paid API or executable document macros."""
from __future__ import annotations
import json
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import threading
import signal
import zipfile
import xml.etree.ElementTree as ET

from telegram_backend import BackendError


def _run(arguments,timeout=60,max_output=2*1024*1024):
    try:
        process=subprocess.Popen(arguments,stdout=subprocess.PIPE,stderr=subprocess.DEVNULL,start_new_session=os.name!='nt')
        try:
            output,_=process.communicate(timeout=timeout)
        except subprocess.TimeoutExpired:
            if os.name!='nt':
                os.killpg(process.pid,signal.SIGKILL)
            else:
                process.kill()
            process.communicate()
            raise
        if process.returncode:
            raise subprocess.CalledProcessError(process.returncode,arguments)
        if len(output)>max_output:
            raise BackendError('media_output_too_large','The media processor produced excessive output.')
        return output
    except FileNotFoundError:
        raise BackendError('processor_unavailable','The required local media/document processor is not installed.') from None
    except (subprocess.SubprocessError,OSError):
        raise BackendError('processing_failed','Unsupported, corrupt or too complex file; local processing did not complete.') from None


def probe_media(path):
    data=json.loads(_run(['ffprobe','-v','error','-protocol_whitelist','file,pipe','-show_entries','format=duration,size:stream=codec_type,codec_name,width,height,duration','-of','json',str(path)],timeout=15))
    result={'duration_seconds':float(data.get('format',{}).get('duration',0) or 0),'streams':[]}
    for stream in data.get('streams',[]):
        result['streams'].append({k:v for k,v in stream.items() if k in ('codec_type','codec_name','width','height','duration')})
        if stream.get('codec_type')=='video':
            result['width']=stream.get('width',0)
            result['height']=stream.get('height',0)
    return result


def image_preview(path,original_limit=4*1024*1024):
    import io
    from PIL import Image,ImageOps,UnidentifiedImageError
    Image.MAX_IMAGE_PIXELS=20_000_000
    try:
        with Image.open(path) as image:
            width,height=image.size
            if width*height>20_000_000:
                raise BackendError('image_too_large','Inline image preview is limited to 20 million pixels; use the original file resource.')
            original_mime={'JPEG':'image/jpeg','PNG':'image/png','WEBP':'image/webp'}.get(image.format)
            if original_mime and Path(path).stat().st_size<=original_limit:
                image.verify()
                return Path(path).read_bytes(),{'preview_mime_type':original_mime,'original_width':width,'original_height':height,
                                             'preview_width':width,'preview_height':height,'original_preserved':True,'preview_uses_original_bytes':True}
            image=ImageOps.exif_transpose(image)
            image.thumbnail((1600,1600))
            image=image.convert('RGB')
            buffer=io.BytesIO()
            image.save(buffer,format='JPEG',quality=90)
            return buffer.getvalue(),{'preview_mime_type':'image/jpeg','original_width':width,'original_height':height,
                                     'preview_width':image.width,'preview_height':image.height,'original_preserved':True,'preview_uses_original_bytes':False}
    except BackendError:
        raise
    except (UnidentifiedImageError,Image.DecompressionBombError,OSError,ValueError):
        raise BackendError('unsupported_image','This image cannot be decoded safely; its original bytes remain available as a resource.') from None


def audio_preview(path):
    info=probe_media(path)
    with tempfile.TemporaryDirectory(prefix='tg-audio-preview-') as temp:
        output=Path(temp)/'preview.wav'
        _run(['ffmpeg','-nostdin','-v','error','-protocol_whitelist','file,pipe','-i',str(path),'-t','60','-ar','16000','-ac','1',str(output)],timeout=45)
        data=output.read_bytes()
    return data,{'preview_seconds':min(info['duration_seconds'],60),'preview_truncated':info['duration_seconds']>60,'original_duration_seconds':info['duration_seconds']}


def video_frames(path,count=4):
    if not 1<=count<=6:
        raise BackendError('invalid_input','Request 1..6 frames.')
    info=probe_media(path)
    if not any(s.get('codec_type')=='video' for s in info['streams']):
        raise BackendError('unsupported_video','This file has no video stream.')
    duration=info['duration_seconds']
    frames=[]
    for index in range(count):
        at=duration*index/count if duration>0 else 0
        data=_run(['ffmpeg','-nostdin','-v','error','-protocol_whitelist','file,pipe','-ss',str(at),'-i',str(path),'-frames:v','1','-vf','scale=640:-2','-f','image2pipe','-vcodec','mjpeg','pipe:1'],timeout=20)
        if data:
            frames.append({'timestamp_seconds':at,'data':data})
    return frames,info


def _zip(path):
    archive=zipfile.ZipFile(path)
    members=archive.infolist()
    if len(members)>10000 or sum(x.file_size for x in members)>80*1024*1024:
        archive.close()
        raise BackendError('document_too_complex','Archive decompression exceeds the bounded document reader limits.')
    return archive


def _xml_text(data):
    if b'<!DOCTYPE' in data.upper() or b'<!ENTITY' in data.upper():
        raise BackendError('unsafe_xml','External entities and DTD declarations are unsupported.')
    root=ET.fromstring(data)
    return '\n'.join(node.text or '' for node in root.iter() if node.tag.rsplit('}',1)[-1] in ('t','p','span','v','text'))


def _legacy(path,filename):
    executable=shutil.which('libreoffice') or shutil.which('soffice')
    if not executable:
        raise BackendError('legacy_converter_unavailable','Legacy DOC/PPT requires the configured LibreOffice converter. Convert to DOCX/PPTX/PDF and upload the converted file as an alternative.')
    with tempfile.TemporaryDirectory(prefix='tg-document-') as temp:
        root=Path(temp)
        source=root/Path(filename).name
        shutil.copyfile(path,source)
        profile=root/'profile'
        profile.mkdir()
        (profile/'user').mkdir()
        (profile/'user'/'registrymodifications.xcu').write_text('''<?xml version="1.0" encoding="UTF-8"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop></item></oor:items>''')
        target='txt:Text' if source.suffix.lower()=='.doc' else 'pdf'
        _run([executable,'--headless','--norestore','--nodefault','--nolockcheck','-env:UserInstallation='+profile.as_uri(),
              '--convert-to',target,'--outdir',str(root),str(source)],timeout=60)
        output=source.with_suffix('.txt' if source.suffix.lower()=='.doc' else '.pdf')
        if not output.is_file():
            raise BackendError('processing_failed','The legacy document could not be converted; no content was returned.')
        if output.suffix=='.txt':
            return output.read_text(encoding='utf-8',errors='replace')
        from pypdf import PdfReader
        return '\n'.join(page.extract_text() or '' for page in PdfReader(output).pages[:100])


def read_document(path,filename,offset=0,max_characters=40000):
    if isinstance(offset,bool) or not isinstance(offset,int) or offset<0 or not 1<=max_characters<=100000:
        raise BackendError('invalid_input','Use a nonnegative character offset and max_characters between 1 and 100000.')
    path=Path(path)
    suffix=Path(filename).suffix.lower()
    if path.stat().st_size>50*1024*1024:
        raise BackendError('document_too_large','Text extraction is limited to 50 MiB; file transfer supports larger originals.')
    try:
        if suffix in ('.txt','.md','.json','.xml','.csv','.log','.html','.htm'):
            content=path.read_bytes()
            if content.startswith(b'\xff\xfe') or content.startswith(b'\xfe\xff'):
                text=content.decode('utf-16')
            else:
                try:
                    text=content.decode('utf-8-sig')
                except UnicodeDecodeError:
                    text=content.decode('cp1251',errors='replace')
        elif suffix=='.pdf':
            from pypdf import PdfReader
            reader=PdfReader(path)
            if reader.is_encrypted:
                raise BackendError('document_encrypted','This PDF requires a password; encrypted content is not extracted.')
            text='\n'.join(f'[Page {i+1}]\n'+(page.extract_text() or '') for i,page in enumerate(reader.pages[:100]))
        elif suffix in ('.docx','.pptx','.ppsx','.odt','.odp','.ods'):
            with _zip(path) as archive:
                if suffix=='.docx':
                    names=['word/document.xml']
                elif suffix in ('.pptx','.ppsx'):
                    import re
                    names=sorted((n for n in archive.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)),key=lambda n:int(re.search(r'slide(\d+)',n).group(1)))
                else:
                    names=['content.xml']
                text='\n'.join(_xml_text(archive.read(name)) for name in names)
        elif suffix in ('.xlsx','.xlsm'):
            from openpyxl import load_workbook
            with _zip(path):
                pass
            workbook=load_workbook(path.open('rb'),read_only=True,data_only=True,keep_links=False)
            pieces=[]
            try:
                for sheet in workbook.worksheets[:30]:
                    pieces.append('[Sheet: '+sheet.title+']')
                    for row in sheet.iter_rows(max_row=1000,max_col=50,values_only=True):
                        pieces.append('\t'.join('' if cell is None else str(cell) for cell in row))
            finally:
                workbook.close()
            text='\n'.join(pieces)
        elif suffix=='.xls':
            import xlrd
            workbook=xlrd.open_workbook(str(path),on_demand=True)
            pieces=[]
            try:
                for sheet in workbook.sheets()[:30]:
                    pieces.append('[Sheet: '+sheet.name+']')
                    pieces.extend('\t'.join(str(x) for x in sheet.row_values(i)[:50]) for i in range(min(sheet.nrows,1000)))
            finally:
                workbook.release_resources()
            text='\n'.join(pieces)
        elif suffix=='.rtf':
            from striprtf.striprtf import rtf_to_text
            text=rtf_to_text(path.read_text(encoding='utf-8',errors='replace'))
        elif suffix in ('.doc','.ppt'):
            text=_legacy(path,filename)
        elif suffix=='.zip':
            with _zip(path) as archive:
                text='\n'.join(f'{item.filename}\t{item.file_size} bytes' for item in archive.infolist())
        elif suffix=='.rar':
            import rarfile
            with rarfile.RarFile(path) as archive:
                text='\n'.join(f'{item.filename}\t{item.file_size} bytes' for item in archive.infolist()[:10000])
        elif suffix=='.7z':
            import py7zr
            with py7zr.SevenZipFile(path,'r') as archive:
                text='\n'.join(archive.getnames()[:10000])
        else:
            raise BackendError('unsupported_document','The original file is available, but this format has no local text extractor.')
        if len(text)>2_000_000:
            text=text[:2_000_000]
        return {'filename':filename,'text':text[offset:offset+max_characters],'next_offset':offset+max_characters if offset+max_characters<len(text) else None,
                'total_characters':len(text),'limits':'100 PDF pages; 30 sheets x 1000 rows x 50 columns; bounded XML/archive reads; no macro execution. Scanned PDFs need OCR.'}
    except BackendError:
        raise
    except Exception:
        raise BackendError('document_read_failed','The document is corrupt, encrypted, unsupported or exceeds reader limits.') from None


_stt_lock=threading.Lock()
_stt_model=None


def transcribe(path,language='ru'):
    global _stt_model
    model_path=os.environ.get('TG_WHISPER_MODEL')
    if (not model_path or not all((Path(model_path)/name).is_file() for name in ('model.bin','config.json','tokenizer.json'))
            or not any((Path(model_path)/name).is_file() for name in ('vocabulary.json','vocabulary.txt'))):
        raise BackendError('transcription_unavailable','A local Whisper model must be provisioned by the server owner; no paid API is called automatically.')
    info=probe_media(path)
    if not 0<info['duration_seconds']<=600:
        raise BackendError('audio_too_long','Transcription supports clips up to 10 minutes.')
    try:
        with _stt_lock:
            if _stt_model is None:
                import onnxruntime as ort
                ort.disable_telemetry_events()
                from faster_whisper import WhisperModel
                _stt_model=WhisperModel(model_path,device='cpu',compute_type='int8',cpu_threads=2,local_files_only=True)
            # Decode with the installed FFmpeg instead of PyAV's unstable open()
            # keyword API; Whisper accepts a mono 16 kHz float32 waveform.
            import numpy as np
            decoded=_run(['ffmpeg','-nostdin','-v','error','-protocol_whitelist','file,pipe','-i',str(path),'-vn','-ar','16000','-ac','1','-f','f32le','pipe:1'],timeout=90,max_output=40*1024*1024)
            audio=np.frombuffer(decoded,dtype='<f4')
            segments,meta=_stt_model.transcribe(audio,language=language or None,beam_size=3,vad_filter=True)
            result=[{'start_seconds':s.start,'end_seconds':s.end,'text':s.text} for s in segments]
        return {'text':' '.join(s['text'].strip() for s in result),'language':meta.language,'segments':result,'provider':'local_whisper','duration_seconds':info['duration_seconds']}
    except BackendError:
        raise
    except Exception:
        raise BackendError('transcription_failed','The local model could not transcribe this audio.') from None


def synthesize(text,language='ru'):
    if not isinstance(text,str) or not 1<=len(text)<=3000 or language not in ('ru','en'):
        raise BackendError('invalid_input','Speech synthesis supports ru/en and at most 3000 characters.')
    with tempfile.TemporaryDirectory(prefix='tg-voice-') as temp:
        path=Path(temp)
        source=path/'speech.txt'
        source.write_text(text,encoding='utf-8')
        wav=path/'speech.wav'
        ogg=path/'speech.ogg'
        _run(['espeak-ng','-b','1','-v',language,'-s','140','-f',str(source),'-w',str(wav)],timeout=60)
        _run(['ffmpeg','-nostdin','-v','error','-i',str(wav),'-c:a','libopus','-b:a','32k',str(ogg)],timeout=60)
        return ogg.read_bytes()
